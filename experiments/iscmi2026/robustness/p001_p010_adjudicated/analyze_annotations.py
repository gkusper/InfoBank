#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from itertools import combinations
from pathlib import Path
from typing import Any, Iterable

import openpyxl


SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parents[3]
BENCHMARK_SCRIPT_DIR = REPO_ROOT / "experiments" / "iscmi2026" / "benchmark"
sys.path.insert(0, str(BENCHMARK_SCRIPT_DIR))

import build_benchmark as bench  # noqa: E402


EXPECTED_PILOTS = [f"P{index:03d}" for index in range(1, 11)]
PASS_SHEETS = ["Annotation1", "Annotation2", "Annotation3"]
FINAL_SHEET = "Finalized_Annotation"
ALLOWED_STATUSES = {"ANNOTATED", "NO_TASK", "SKIP_UNCERTAIN", "PENDING"}
CORE_REQUIRED_ANNOTATED_FIELDS = {
    "task_id",
    "message_id",
    "chronological_rank",
    "transition",
    "resulting_state",
    "confidence",
    "annotation_status",
}


def json_text(value: Any, *, compact: bool = False) -> str:
    if compact:
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n"


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json_text(value), encoding="utf-8", newline="\n")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def value_present(value: Any) -> bool:
    return value is not None and value != ""


def normalize_cell(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def load_sheet(workbook_path: Path, sheet_name: str) -> tuple[list[str], list[dict[str, Any]]]:
    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Workbook does not contain required sheet {sheet_name!r}")
    ws = wb[sheet_name]
    headers = [
        bench.HEADER_ALIASES.get(str(cell.value).strip(), str(cell.value).strip())
        if cell.value is not None
        else ""
        for cell in ws[1]
    ]
    missing = [column for column in bench.REQUIRED_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"{sheet_name} is missing required columns: {missing}")
    rows: list[dict[str, Any]] = []
    for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        padded = list(values) + [""] * (len(headers) - len(values))
        if not any(value_present(normalize_cell(value)) for value in padded[: len(headers)]):
            continue
        record = {
            header: normalize_cell(padded[index]) if index < len(padded) else ""
            for index, header in enumerate(headers)
        }
        record["_source_workbook_row"] = row_number
        for field in ("chronological_rank", "confidence"):
            if value_present(record.get(field)):
                number = float(record[field])
                if not number.is_integer():
                    raise ValueError(f"{field} must be an integer in {sheet_name} row {row_number}")
                record[field] = int(number)
        for field in bench.DATE_FIELDS:
            value = record.get(field)
            if isinstance(value, (int, float)):
                record[field] = (datetime(1899, 12, 30) + timedelta(days=float(value))).date().isoformat()
        rows.append(record)
    return headers, rows


def load_workbook_sheets(workbook_path: Path) -> tuple[list[str], dict[str, list[str]], dict[str, list[dict[str, Any]]]]:
    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    sheet_names = list(wb.sheetnames)
    rows_by_sheet: dict[str, list[dict[str, Any]]] = {}
    annotators_by_sheet: dict[str, list[str]] = {}
    for sheet_name in sheet_names:
        _headers, rows = load_sheet(workbook_path, sheet_name)
        rows_by_sheet[sheet_name] = rows
        annotators_by_sheet[sheet_name] = sorted(
            {str(row.get("annotator_id")).strip() for row in rows if value_present(row.get("annotator_id"))}
        )
    return sheet_names, annotators_by_sheet, rows_by_sheet


def message_key(row: dict[str, Any]) -> tuple[str, str]:
    return str(row["pilot_id"]), str(row["message_id"])


def status_counts(rows: list[dict[str, Any]]) -> dict[str, int]:
    return dict(sorted(Counter(str(row.get("annotation_status") or "") for row in rows).items()))


def annotated_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [row for row in rows if row.get("annotation_status") == bench.ANNOTATED]


def task_snapshots(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in annotated_rows(rows):
        grouped[(str(row["pilot_id"]), str(row["task_id"]))].append(row)
    snapshots = []
    for (pilot_id, task_id), task_rows in sorted(grouped.items(), key=lambda item: (item[0][0], bench.task_sort_key(item[0][1]))):
        ordered = sorted(task_rows, key=lambda row: (int(row["chronological_rank"]), int(row["_source_workbook_row"])))
        final = ordered[-1]
        snapshots.append(
            {
                "pilot_id": pilot_id,
                "task_id": task_id,
                "final_state": final.get("resulting_state"),
                "last_transition": final.get("transition"),
                "requester": final.get("requester") or None,
                "actor": final.get("actor") or None,
                "evidence_message_ids": list(dict.fromkeys(str(row["message_id"]) for row in ordered)),
            }
        )
    return snapshots


def final_state_counts(rows: list[dict[str, Any]]) -> dict[str, int]:
    return dict(sorted(Counter(str(row["final_state"]) for row in task_snapshots(rows)).items()))


def message_presence(rows: list[dict[str, Any]]) -> dict[tuple[str, str], bool]:
    result: dict[tuple[str, str], bool] = defaultdict(bool)
    for row in rows:
        key = message_key(row)
        result[key] = bool(result[key] or row.get("annotation_status") == bench.ANNOTATED)
    return dict(result)


def message_structure(rows: list[dict[str, Any]]) -> dict[tuple[str, str], tuple[tuple[str, str], ...]]:
    result: dict[tuple[str, str], list[tuple[str, str]]] = defaultdict(list)
    for row in annotated_rows(rows):
        result[message_key(row)].append((str(row.get("transition") or ""), str(row.get("resulting_state") or "")))
    return {key: tuple(sorted(values)) for key, values in result.items()}


def message_task_counts(rows: list[dict[str, Any]]) -> dict[tuple[str, str], int]:
    counts: dict[tuple[str, str], int] = defaultdict(int)
    for row in annotated_rows(rows):
        counts[message_key(row)] += 1
    return dict(counts)


def continuity_edges(rows: list[dict[str, Any]]) -> set[tuple[str, str, str]]:
    by_task: dict[tuple[str, str], set[str]] = defaultdict(set)
    for row in annotated_rows(rows):
        by_task[(str(row["pilot_id"]), str(row["task_id"]))].add(str(row["message_id"]))
    edges: set[tuple[str, str, str]] = set()
    for (pilot_id, _task_id), message_ids in by_task.items():
        for left, right in combinations(sorted(message_ids), 2):
            edges.add((pilot_id, left, right))
    return edges


def f1_counts(left: set[Any], right: set[Any]) -> dict[str, float | int]:
    overlap = len(left & right)
    precision = overlap / len(left) if left else (1.0 if not right else 0.0)
    recall = overlap / len(right) if right else (1.0 if not left else 0.0)
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    union = len(left | right)
    return {
        "left_count": len(left),
        "right_count": len(right),
        "overlap": overlap,
        "jaccard": round(overlap / union, 8) if union else 1.0,
        "precision_left_to_right": round(precision, 8),
        "recall_left_to_right": round(recall, 8),
        "f1": round(f1, 8),
    }


def pairwise_agreement(sheet_rows: dict[str, list[dict[str, Any]]], message_ids: list[tuple[str, str]]) -> list[dict[str, Any]]:
    rows = []
    for left, right in combinations(PASS_SHEETS, 2):
        left_presence = message_presence(sheet_rows[left])
        right_presence = message_presence(sheet_rows[right])
        presence_matches = sum(
            int(bool(left_presence.get(key, False)) == bool(right_presence.get(key, False)))
            for key in message_ids
        )
        left_structure = message_structure(sheet_rows[left])
        right_structure = message_structure(sheet_rows[right])
        structure_matches = sum(
            int(left_structure.get(key, tuple()) == right_structure.get(key, tuple()))
            for key in message_ids
        )
        left_edges = continuity_edges(sheet_rows[left])
        right_edges = continuity_edges(sheet_rows[right])
        rows.append(
            {
                "left": left,
                "right": right,
                "message_count": len(message_ids),
                "task_presence_agreement": round(presence_matches / len(message_ids), 8),
                "transition_state_structure_agreement": round(structure_matches / len(message_ids), 8),
                "continuity_edge_agreement": f1_counts(left_edges, right_edges),
            }
        )
    return rows


def describe_structure(values: tuple[tuple[str, str], ...]) -> str:
    if not values:
        return "NONE"
    counts = Counter(values)
    return "; ".join(f"{count}x {transition}/{state}" for (transition, state), count in sorted(counts.items()))


def disagreement_rows(sheet_rows: dict[str, list[dict[str, Any]]], message_ids: list[tuple[str, str]]) -> list[dict[str, Any]]:
    structures = {sheet: message_structure(rows) for sheet, rows in sheet_rows.items() if sheet in PASS_SHEETS}
    counts = {sheet: message_task_counts(rows) for sheet, rows in sheet_rows.items() if sheet in PASS_SHEETS}
    disagreements = []
    for pilot_id, message_id in message_ids:
        key = (pilot_id, message_id)
        signatures = {
            sheet: {
                "task_mentions": counts[sheet].get(key, 0),
                "transition_state": describe_structure(structures[sheet].get(key, tuple())),
            }
            for sheet in PASS_SHEETS
        }
        if len({json_text(value, compact=True) for value in signatures.values()}) > 1:
            rank_match = re.search(r"::turn_(\d+)$", message_id)
            rank = int(rank_match.group(1)) + 1 if rank_match else None
            disagreements.append(
                {
                    "pilot_id": pilot_id,
                    "chronological_rank": rank,
                    "message_id": message_id,
                    "passes": signatures,
                }
            )
    return disagreements


def validate_finalized(
    headers: list[str],
    rows: list[dict[str, Any]],
    message_by_id: dict[str, dict[str, Any]],
    thread_by_pilot: dict[str, str],
) -> dict[str, Any]:
    issues: list[str] = []
    bench.validate_annotation_rows(headers, rows, message_by_id, thread_by_pilot, EXPECTED_PILOTS)
    if any(row.get("annotation_status") == "PENDING" for row in rows):
        issues.append("PENDING annotation remains in finalized sheet")
    duplicate_keys = Counter(
        (
            row.get("pilot_id"),
            row.get("task_id"),
            row.get("message_id"),
            row.get("transition"),
            row.get("resulting_state"),
        )
        for row in annotated_rows(rows)
    )
    duplicates = [
        {"key": list(key), "count": count}
        for key, count in duplicate_keys.items()
        if count > 1
    ]
    for row in rows:
        workbook_row = row["_source_workbook_row"]
        status = str(row.get("annotation_status") or "")
        if status not in ALLOWED_STATUSES:
            issues.append(f"row {workbook_row}: invalid annotation_status {status!r}")
            continue
        if status == bench.ANNOTATED:
            for field in CORE_REQUIRED_ANNOTATED_FIELDS:
                if not value_present(row.get(field)):
                    issues.append(f"row {workbook_row}: ANNOTATED row missing {field}")
            if row.get("transition") not in bench.ALLOWED_TRANSITIONS:
                issues.append(f"row {workbook_row}: invalid transition {row.get('transition')!r}")
            if row.get("resulting_state") not in bench.ALLOWED_STATES:
                issues.append(f"row {workbook_row}: invalid resulting_state {row.get('resulting_state')!r}")
            if row.get("resulting_state") == "CLOSED" and not value_present(row.get("closure_reason")):
                issues.append(f"row {workbook_row}: CLOSED row lacks closure_reason")
            if value_present(row.get("closure_reason")) and row.get("resulting_state") != "CLOSED":
                issues.append(f"row {workbook_row}: non-CLOSED row has closure_reason")
            confidence = row.get("confidence")
            if not isinstance(confidence, int) or not 1 <= confidence <= 5:
                issues.append(f"row {workbook_row}: confidence is not an integer 1-5")
        elif status in {"NO_TASK", "SKIP_UNCERTAIN"}:
            semantic_fields = [
                "task_id",
                "transition",
                "resulting_state",
                "closure_reason",
                "waiting_for",
                "blocked_by",
            ]
            present = [field for field in semantic_fields if value_present(row.get(field))]
            if present:
                issues.append(f"row {workbook_row}: {status} row has semantic fields {present}")
    if duplicates:
        issues.append(f"duplicate semantic rows found: {duplicates[:5]}")
    by_pilot: dict[str, list[int]] = defaultdict(list)
    for row in rows:
        by_pilot[str(row["pilot_id"])].append(int(row["chronological_rank"]))
        expected = message_by_id[str(row["message_id"])]["chronological_rank"]
        if int(row["chronological_rank"]) != expected:
            issues.append(f"row {row['_source_workbook_row']}: chronological rank does not match packet")
    for pilot_id, ranks in by_pilot.items():
        if min(ranks) < 1:
            issues.append(f"{pilot_id}: chronological ranks must be positive")
    return {
        "status": "PASS" if not issues else "FAIL",
        "issues": issues,
        "duplicate_semantic_rows": duplicates,
        "mechanical_corrections": [],
    }


def sheet_summary(rows: list[dict[str, Any]]) -> dict[str, Any]:
    snapshots = task_snapshots(rows)
    return {
        "rows": len(rows),
        "annotated_task_event_rows": len(annotated_rows(rows)),
        "messages_with_any_row": len({row["message_id"] for row in rows}),
        "messages_with_task_presence": len({row["message_id"] for row in annotated_rows(rows)}),
        "final_tasks": len(snapshots),
        "final_state_distribution": dict(sorted(Counter(row["final_state"] for row in snapshots).items())),
        "annotation_status_distribution": status_counts(rows),
        "transition_distribution": dict(sorted(Counter(row["transition"] for row in annotated_rows(rows)).items())),
        "confidence_distribution": dict(sorted(Counter(str(row["confidence"]) for row in annotated_rows(rows)).items())),
    }


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def old_vs_clean(clean_benchmark_dir: Path) -> dict[str, Any] | None:
    if not (clean_benchmark_dir / "benchmark_statistics.json").exists():
        return None
    old_snapshots = [
        row for row in read_jsonl(REPO_ROOT / "experiments" / "iscmi2026" / "benchmark" / "task_snapshots.jsonl")
        if row["pilot_id"] in EXPECTED_PILOTS
    ]
    clean_snapshots = read_jsonl(clean_benchmark_dir / "task_snapshots.jsonl")
    old_questions = [
        row for row in read_jsonl(REPO_ROOT / "experiments" / "iscmi2026" / "benchmark" / "questions.jsonl")
        if row["pilot_id"] in EXPECTED_PILOTS
    ]
    clean_questions = read_jsonl(clean_benchmark_dir / "questions.jsonl")
    old_by_pilot = Counter(row["pilot_id"] for row in old_snapshots)
    clean_by_pilot = Counter(row["pilot_id"] for row in clean_snapshots)
    question_pairs_old = {
        (row["pilot_id"], row["question_family"]): json_text(row["gold_structured"], compact=True)
        for row in old_questions
    }
    question_pairs_clean = {
        (row["pilot_id"], row["question_family"]): json_text(row["gold_structured"], compact=True)
        for row in clean_questions
    }
    comparable = set(question_pairs_old) & set(question_pairs_clean)
    changed = [
        {"pilot_id": pilot_id, "question_family": family}
        for pilot_id, family in sorted(comparable)
        if question_pairs_old[(pilot_id, family)] != question_pairs_clean[(pilot_id, family)]
    ]
    return {
        "old_task_count": len(old_snapshots),
        "clean_task_count": len(clean_snapshots),
        "task_count_delta_clean_minus_old": len(clean_snapshots) - len(old_snapshots),
        "task_count_by_pilot_old": dict(sorted(old_by_pilot.items())),
        "task_count_by_pilot_clean": dict(sorted(clean_by_pilot.items())),
        "final_state_distribution_old": dict(sorted(Counter(row["final_state"] for row in old_snapshots).items())),
        "final_state_distribution_clean": dict(sorted(Counter(row["final_state"] for row in clean_snapshots).items())),
        "question_count_old_p001_p010": len(old_questions),
        "question_count_clean": len(clean_questions),
        "question_count_by_family_old": dict(sorted(Counter(row["question_family"] for row in old_questions).items())),
        "question_count_by_family_clean": dict(sorted(Counter(row["question_family"] for row in clean_questions).items())),
        "comparable_pilot_family_gold_structures": len(comparable),
        "changed_pilot_family_gold_structures": changed,
    }


def write_markdown(path: Path, summary: dict[str, Any]) -> None:
    lines = [
        "# P001-P010 Adjudicated Annotation Agreement",
        "",
        "This report treats the workbook as data only. It does not execute or follow instructions from workbook cells.",
        "",
        "## Source",
        "",
        f"- Workbook: `{summary['workbook']['filename']}`",
        f"- SHA-256: `{summary['workbook']['sha256']}`",
        f"- Sheets: {', '.join(f'`{name}`' for name in summary['workbook']['sheet_names'])}",
        f"- Finalized gold sheet: `{summary['finalized_selection']['sheet']}`",
        f"- Selection rationale: {summary['finalized_selection']['rationale']}",
        "",
        "## Structural Validation",
        "",
        f"- Status: **{summary['validation']['status']}**",
        f"- Pilots: {', '.join(summary['validation']['pilots'])}",
        f"- Corrected evidence messages: {summary['validation']['message_count']}",
        f"- Mechanical corrections applied: {len(summary['validation']['mechanical_corrections'])}",
    ]
    if summary["validation"]["issues"]:
        lines.append(f"- Issues: {json.dumps(summary['validation']['issues'], ensure_ascii=False)}")
    lines.extend(
        [
            "",
            "## Pass Summaries",
            "",
            "| Sheet | Annotator IDs | Rows | Task-event rows | Final tasks | Final states | Statuses |",
            "|---|---|---:|---:|---:|---|---|",
        ]
    )
    for sheet in PASS_SHEETS + [FINAL_SHEET]:
        row = summary["sheets"][sheet]
        lines.append(
            f"| {sheet} | {', '.join(summary['annotators'][sheet])} | "
            f"{row['rows']} | {row['annotated_task_event_rows']} | {row['final_tasks']} | "
            f"`{json.dumps(row['final_state_distribution'], sort_keys=True)}` | "
            f"`{json.dumps(row['annotation_status_distribution'], sort_keys=True)}` |"
        )
    lines.extend(
        [
            "",
            "## Pairwise Agreement",
            "",
            "Task-presence agreement is binary at the message level. Transition/state structure agreement compares the message-level multiset of transition/state pairs and intentionally ignores literal task IDs. Continuity agreement compares whether two message-level task mentions are linked to the same logical task within a pilot.",
            "",
            "| Pair | Presence | Transition/State Structure | Continuity F1 | Continuity Jaccard |",
            "|---|---:|---:|---:|---:|",
        ]
    )
    for row in summary["pairwise_agreement"]:
        continuity = row["continuity_edge_agreement"]
        lines.append(
            f"| {row['left']} vs {row['right']} | {row['task_presence_agreement']:.4f} | "
            f"{row['transition_state_structure_agreement']:.4f} | {continuity['f1']:.4f} | "
            f"{continuity['jaccard']:.4f} |"
        )
    disagreements = summary["decomposition_or_structure_disagreements"]
    lines.extend(
        [
            "",
            "## Disagreements",
            "",
            f"Messages with task decomposition or transition/state interpretation differences across the three annotation passes: {len(disagreements)}.",
            "",
        ]
    )
    for item in disagreements:
        lines.append(
            f"- `{item['pilot_id']}` rank {item['chronological_rank']} `{item['message_id']}`: "
            + "; ".join(
                f"{sheet}={data['task_mentions']} mentions, {data['transition_state']}"
                for sheet, data in item["passes"].items()
            )
        )
    if summary.get("old_vs_clean"):
        change = summary["old_vs_clean"]
        lines.extend(
            [
                "",
                "## Original-vs-Adjudicated Sensitivity",
                "",
                f"- Original P001-P010 task count: {change['old_task_count']}",
                f"- Adjudicated task count: {change['clean_task_count']}",
                f"- Task-count delta: {change['task_count_delta_clean_minus_old']}",
                f"- Original final states: `{json.dumps(change['final_state_distribution_old'], sort_keys=True)}`",
                f"- Adjudicated final states: `{json.dumps(change['final_state_distribution_clean'], sort_keys=True)}`",
                f"- Original P001-P010 question count: {change['question_count_old_p001_p010']}",
                f"- Adjudicated question count: {change['question_count_clean']}",
                f"- Changed comparable pilot/family gold structures: {len(change['changed_pilot_family_gold_structures'])} of {change['comparable_pilot_family_gold_structures']}",
                "",
                "This is an adjudicated annotation-quality robustness subset, not a corpus-wide annotation-error estimate.",
            ]
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8", newline="\n")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Analyze P001-P010 adjudicated MailEx annotation passes.")
    parser.add_argument("--workbook", type=Path, required=True, help="Attached adjudicated workbook.")
    parser.add_argument("--packet-dir", type=Path, required=True, help="Corrected external MailEx packet directory.")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=SCRIPT_DIR,
        help="Output directory for annotation_agreement_report.md and annotation_agreement_summary.json.",
    )
    parser.add_argument(
        "--clean-benchmark-dir",
        type=Path,
        default=SCRIPT_DIR / "benchmark",
        help="Optional clean benchmark directory for original-vs-adjudicated sensitivity.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    sheet_names, annotators, sheet_rows = load_workbook_sheets(args.workbook)
    required = set(PASS_SHEETS + [FINAL_SHEET])
    missing = sorted(required - set(sheet_names))
    if missing:
        raise ValueError(f"Workbook is missing required annotation sheets: {missing}")

    headers, final_rows = load_sheet(args.workbook, FINAL_SHEET)
    message_index, thread_by_pilot, _packet_hash, _pilot_hashes, _snippets = bench.read_pilot_packet(
        args.packet_dir,
        EXPECTED_PILOTS,
        require_exact_files=False,
    )
    message_by_id = {row["message_id"]: row for row in message_index}
    validation = validate_finalized(headers, final_rows, message_by_id, thread_by_pilot)
    if validation["status"] != "PASS":
        raise ValueError(f"Finalized annotation validation failed: {validation['issues']}")
    message_ids = sorted(
        {(row["pilot_id"], row["message_id"]) for row in message_index},
        key=lambda key: (key[0], message_by_id[key[1]]["chronological_rank"]),
    )
    summary = {
        "workbook": {
            "filename": args.workbook.name,
            "sha256": sha256_file(args.workbook),
            "sheet_names": sheet_names,
        },
        "annotators": {sheet: annotators.get(sheet, []) for sheet in PASS_SHEETS + [FINAL_SHEET]},
        "finalized_selection": {
            "sheet": FINAL_SHEET,
            "rationale": (
                "The workbook contains an explicit Finalized_Annotation worksheet with the same "
                "P001-P010 schema and message coverage as the annotation passes. Its annotator "
                "provenance is KG_A15, indicating a corrected/final adjudication layer rather "
                "than a fourth independent vote."
            ),
        },
        "validation": {
            **validation,
            "pilots": EXPECTED_PILOTS,
            "message_count": len(message_index),
        },
        "sheets": {sheet: sheet_summary(sheet_rows[sheet]) for sheet in PASS_SHEETS + [FINAL_SHEET]},
        "pairwise_agreement": pairwise_agreement(sheet_rows, message_ids),
        "decomposition_or_structure_disagreements": disagreement_rows(sheet_rows, message_ids),
        "old_vs_clean": old_vs_clean(args.clean_benchmark_dir),
    }
    write_json(args.output_dir / "annotation_agreement_summary.json", summary)
    write_markdown(args.output_dir / "annotation_agreement_report.md", summary)
    print(
        json_text(
            {
                "report": str(args.output_dir / "annotation_agreement_report.md"),
                "summary": str(args.output_dir / "annotation_agreement_summary.json"),
                "finalized_sheet": FINAL_SHEET,
                "messages": len(message_index),
                "final_tasks": summary["sheets"][FINAL_SHEET]["final_tasks"],
                "validation": validation["status"],
            },
            compact=True,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
